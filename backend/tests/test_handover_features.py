"""Features added for handover: card payments, sold out, quick sale, comps, modifiers,
combos, happy hour, sections, Excel export, the emailed closing and menu pictures."""

import base64
import io
from datetime import datetime, time, timedelta
from decimal import Decimal
from zoneinfo import ZoneInfo

import pytest
from conftest import login
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.main import create_app
from app.menu.models import PriceRule
from app.menu.service import rule_is_on
from app.printing.models import PrinterName
from app.reports.mailer import send_pending
from app.reports.models import EmailOutbox

SAM, JANE = "8642", "1357"
NAIROBI = ZoneInfo("Africa/Nairobi")


@pytest.fixture
def people(add_staff):
    add_staff("Sam", "supervisor", SAM)
    add_staff("Jane", "waiter", JANE)


@pytest.fixture
def m(client, owner, menu):
    """Menu lookups: options by display name, items by name, modifier options by label."""
    data = client.get("/api/menu", headers=owner).json()
    opts, items, mods = {}, {}, {}
    for cat in data["categories"]:
        for item in cat["items"]:
            items[item["name"]] = item
            for o in item["options"]:
                opts[f"{item['name']} - {o['label']}" if o["label"] else item["name"]] = o["id"]
            for g in item["modifier_groups"]:
                for o in g["options"]:
                    mods[o["label"]] = o["id"]
    return {"opts": opts, "items": items, "mods": mods}


def stock_level(client, headers, name):
    return Decimal(next(s["level"] for s in client.get("/api/stock", headers=headers).json() if s["name"] == name))


def deliver(client, headers, name, units):
    sid = next(s["id"] for s in client.get("/api/stock", headers=headers).json() if s["name"] == name)
    client.post("/api/stock/deliveries", headers=headers, json={"lines": [{"stock_item_id": sid, "units": units}]})


def tab(client, headers, name="Kamau"):
    return client.post("/api/bills", json={"tab_name": name}, headers=headers).json()


def send(client, headers, bill, *items):
    body = {"items": [dict(sell_option_id=o, quantity=q, modifier_option_ids=mods) for o, q, mods in items]}
    return client.post(f"/api/bills/{bill['id']}/send", json=body, headers=headers)


def test_card_payments_count_as_non_cash(client, owner, people, m, printers):
    jane = login(client, JANE)
    bill = send(client, jane, tab(client, jane), (m["opts"]["Tusker"], 2, [])).json()
    pay = lambda **b: client.post(f"/api/bills/{bill['id']}/payments", headers=jane, json=b)  # noqa: E731
    assert pay(method="card", amount=600).status_code == 422
    r = pay(method="card", amount=600, card_ref="8812")
    assert r.status_code == 201 and r.json()["bill"]["payments"][0]["card_ref"] == "8812"
    assert "Card ref 8812" in printers[PrinterName.counter].printed[-1]
    c = client.get("/api/shifts/me", headers=jane).json()["cash_up"]
    assert (c["total_sales"], c["card"], c["cash_to_hand_over"]) == (600, 600, 0)
    assert c["mpesa_payments"][0] == {**c["mpesa_payments"][0], "method": "card", "ref": "8812"}
    assert client.get("/api/closing", headers=owner).json()["report"]["sales"]["card"] == 600


def test_sold_out_items_cannot_be_sent(client, owner, people, m):
    sam, jane = login(client, SAM), login(client, JANE)
    tusker = m["items"]["Tusker"]
    url = f"/api/menu/items/{tusker['id']}/availability"
    assert client.post(url, json={"available": False}, headers=jane).status_code == 403
    assert client.post(url, json={"available": False}, headers=sam).status_code == 204
    assert next(i for c in client.get("/api/menu", headers=jane).json()["categories"] for i in c["items"]
                if i["name"] == "Tusker")["available"] is False
    r = send(client, jane, tab(client, jane), (m["opts"]["Tusker"], 1, []))
    assert r.status_code == 409 and "sold out" in r.json()["detail"]
    client.post(url, json={"available": True}, headers=sam)
    assert send(client, jane, tab(client, jane, "Otieno"), (m["opts"]["Tusker"], 1, [])).status_code == 200


def test_an_order_needs_nothing_but_its_number(client, owner, m, printers):
    bill = client.post("/api/bills", json={}, headers=owner).json()
    assert bill["place"] == f"Order #{bill['number']:04d}" and bill["section"] == "Bar"
    send(client, owner, bill, (m["opts"]["Tusker"], 1, []))
    assert f"ORDER #{bill['number']:04d}" in printers[PrinterName.counter].printed[-1]  # the runner finds it by number
    quick = client.post("/api/bills", json={"quick": True}, headers=owner).json()  # older screens
    assert quick["place"] == f"Order #{quick['number']:04d}"
    # Tapping NEW ORDER again while that one is still empty gives the same order back.
    assert client.post("/api/bills", json={}, headers=owner).json()["id"] == quick["id"]
    assert client.post("/api/bills", json={"quick": True, "tab_name": "x"}, headers=owner).status_code == 422


def test_a_typed_table_number(client, owner, m, people):
    floor = client.get("/api/floor", headers=owner).json()["tables"]
    t7 = next(t for t in floor if t["name"] == "T7")
    # "7" is a table that's set up: the order sits there, with its section.
    seven = client.post("/api/bills", json={"table": "7"}, headers=owner).json()
    assert seven["table_id"] == t7["id"] and seven["place"] == "T7" and seven["section"] == t7["section"]
    jane = login(client, JANE)
    r = client.post("/api/bills", json={"table": "t7"}, headers=jane)
    assert r.status_code == 409 and "Njeri" in r.json()["detail"]
    # "S12" isn't set up (a table brought out from the stall): it's used as typed.
    stall = client.post("/api/bills", json={"table": "S12"}, headers=jane).json()
    assert stall["place"] == "Table S12" and stall["table_id"] is None and stall["table_label"] == "S12"
    assert client.post("/api/bills", json={"table": "<b>"}, headers=jane).status_code == 422
    # An order by number gets a table later.
    order = client.post("/api/bills", json={}, headers=jane).json()
    moved = client.post(f"/api/bills/{order['id']}/move", json={"table": "57"}, headers=jane).json()
    assert moved["place"] == "Table 57" and moved["number"] == order["number"]
    moved = client.post(f"/api/bills/{order['id']}/move", json={"table": "T8"}, headers=jane).json()
    assert moved["place"] == "T8" and moved["table_label"] is None
    mine = [b for b in client.get("/api/floor", headers=jane).json()["bills"] if b["mine"]]
    assert {b["place"] for b in mine} == {"Table S12", "T8"}


def test_comps_are_free_but_still_leave_stock(client, owner, people, m, printers):
    jane = login(client, JANE)
    deliver(client, owner, "Guinness 500ml", 10)
    bill = send(client, jane, tab(client, jane), (m["opts"]["Guinness"], 2, [])).json()
    line = bill["lines"][0]
    url = f"/api/bills/{bill['id']}/lines/{line['id']}/comp"
    assert client.post(url, json={"reason": "staff_drink"}, headers=jane).status_code == 403
    approval = client.post("/api/auth/approvals", headers=jane,
                           json={"action": "comp_item", "approver_pin": SAM, "reason": "Staff drink"}).json()["id"]
    r = client.post(url, json={"reason": "staff_drink", "approval_id": approval}, headers=jane)
    assert r.status_code == 200 and r.json()["total"] == 320
    assert stock_level(client, owner, "Guinness 500ml") == 8
    client.post(f"/api/bills/{bill['id']}/payments", headers=jane, json={"method": "cash", "amount": 320})
    assert "Staff drink (free)" in printers[PrinterName.counter].printed[-1]
    comps = client.get("/api/closing", headers=owner).json()["report"]["comps"]
    assert comps["total"] == 320 and comps["items"][0]["approved_by"] == "Sam"


def test_modifiers_are_checked_priced_printed_and_take_stock(client, owner, m, printers):
    deliver(client, owner, "Coke 500ml", 5)
    nyama, jw = m["opts"]["Nyama choma 1/2 kg"], m["opts"]["Johnnie Walker Red - Single tot"]
    bill = tab(client, owner)
    r = send(client, owner, bill, (nyama, 1, []))
    assert r.status_code == 422 and "side" in r.json()["detail"]
    r = send(client, owner, bill, (nyama, 1, [m["mods"]["Ugali"], m["mods"]["Chips"]]))
    assert r.status_code == 422 and "at most 1" in r.json()["detail"]
    r = send(client, owner, bill, (nyama, 1, [m["mods"]["Ugali"], m["mods"]["Well done"]]), (jw, 1, [m["mods"]["Coke"]]))
    assert r.status_code == 200
    assert r.json()["total"] == 800 + 250 + 100
    assert "+ Ugali" in printers[PrinterName.kitchen].printed[-1] and "+ Well done" in printers[PrinterName.kitchen].printed[-1]
    assert "+ Coke  100" in printers[PrinterName.counter].printed[-1]
    assert stock_level(client, owner, "Coke 500ml") == 4
    assert stock_level(client, owner, "Maize flour") == -250
    assert send(client, owner, bill, (m["opts"]["Tusker"], 1, [m["mods"]["Coke"]])).status_code == 422


def test_combo_sends_parts_and_charges_once(client, owner, m, printers):
    deliver(client, owner, "Coke 500ml", 5)
    deliver(client, owner, "Chicken (quarter)", 5)
    bill = send(client, owner, tab(client, owner), (m["opts"]["Chicken & chips + Coke"], 1, [])).json()
    assert bill["total"] == 750
    parent = next(l for l in bill["lines"] if l["combo_parent_id"] is None)
    parts = [l for l in bill["lines"] if l["combo_parent_id"] == parent["id"]]
    assert sorted(p["item_name"] for p in parts) == ["Chicken & chips", "Coke"]
    assert "1 x Chicken & chips" in printers[PrinterName.kitchen].printed[-1]
    assert "1 x Coke" in printers[PrinterName.counter].printed[-1]
    assert stock_level(client, owner, "Coke 500ml") == 4 and stock_level(client, owner, "Chicken (quarter)") == 4

    url = lambda line: f"/api/bills/{bill['id']}/lines/{line['id']}/void"  # noqa: E731
    assert client.post(url(parts[0]), json={"reason": "not_prepared"}, headers=owner).status_code == 409
    r = client.post(url(parent), json={"reason": "not_prepared"}, headers=owner)
    assert r.json()["total"] == 0 and all(l["voided"] for l in r.json()["lines"])
    assert stock_level(client, owner, "Coke 500ml") == 5

    bill = send(client, owner, tab(client, owner, "Wanjiru"), (m["opts"]["Chicken & chips + Coke"], 1, [])).json()
    client.post(f"/api/bills/{bill['id']}/payments", headers=owner, json={"method": "cash", "amount": 750})
    receipt = printers[PrinterName.counter].printed[-1]
    assert "1 x Chicken & chips + Coke" in receipt and "x Coke " not in receipt.replace("+ Coke", "")


def test_happy_hour_prices_at_order_time(client, owner, m, db, business, printers):
    beers = next(c for c in client.get("/api/menu", headers=owner).json()["categories"] if c["name"] == "Beers")["id"]
    for r in db.scalars(select(PriceRule)):
        db.delete(r)
    db.commit()
    now = datetime.now(NAIROBI)
    start, end = (now - timedelta(minutes=30)).time(), (now + timedelta(minutes=30)).time()
    r = client.post("/api/menu/price-rules", headers=owner, json={
        "name": "Happy hour", "percent_off": 20, "start_time": start.strftime("%H:%M"), "end_time": end.strftime("%H:%M"),
        "category_ids": [beers]})
    assert r.status_code == 201
    menu = client.get("/api/menu", headers=owner).json()
    assert menu["happy_hours"][0]["percent_off"] == 20
    tusker = next(i for c in menu["categories"] for i in c["items"] if i["name"] == "Tusker")["options"][0]
    assert (tusker["price"], tusker["offer_price"]) == (300, 240)
    bill = send(client, owner, tab(client, owner), (m["opts"]["Tusker"], 2, []), (m["opts"]["Coke"], 1, [])).json()
    assert bill["total"] == 480 + 100
    client.post(f"/api/bills/{bill['id']}/payments", headers=owner, json={"method": "cash", "amount": 580})
    assert "Happy hour (was 300)" in printers[PrinterName.counter].printed[-1]
    bad = client.post("/api/menu/price-rules", headers=owner, json={"percent_off": 10, "start_time": "17:00", "end_time": "17:00"})
    assert bad.status_code == 422


def test_happy_hour_window_across_midnight():
    rule = PriceRule(name="Late", percent_off=10, days=[4], start_time=time(22, 0), end_time=time(1, 0))  # Friday night
    friday_late = datetime(2026, 10, 16, 23, 30, tzinfo=NAIROBI)
    saturday_early = datetime(2026, 10, 17, 0, 30, tzinfo=NAIROBI)
    saturday_late = datetime(2026, 10, 17, 23, 30, tzinfo=NAIROBI)
    assert rule_is_on(rule, friday_late, NAIROBI) and rule_is_on(rule, saturday_early, NAIROBI)
    assert not rule_is_on(rule, saturday_late, NAIROBI)


def test_sections_and_sales_by_section(client, owner, m):
    tables = {t["name"]: t for t in client.get("/api/floor", headers=owner).json()["tables"]}
    assert tables["Garden 1"]["section"] == "Garden"
    garden = client.post("/api/bills", json={"table_id": tables["Garden 1"]["id"]}, headers=owner).json()
    assert garden["section"] == "Garden"
    send(client, owner, garden, (m["opts"]["Tusker"], 1, []))
    client.post(f"/api/bills/{garden['id']}/payments", headers=owner, json={"method": "cash", "amount": 300})
    bar = send(client, owner, tab(client, owner), (m["opts"]["Coke"], 1, [])).json()
    client.post(f"/api/bills/{bar['id']}/payments", headers=owner, json={"method": "cash", "amount": 100})
    sections = {s["section"]: s["sales"] for s in client.get("/api/closing", headers=owner).json()["report"]["sections"]}
    assert sections == {"Garden": 300, "Bar": 100}
    r = client.post("/api/tables", headers=owner, json={"name": "VIP 1", "section": "VIP"})
    assert r.status_code == 201 and r.json()["section"] == "VIP"


def test_excel_exports(client, owner, people, m):
    bill = send(client, owner, tab(client, owner), (m["opts"]["Tusker"], 1, [])).json()
    client.post(f"/api/bills/{bill['id']}/payments", headers=owner, json={"method": "cash", "amount": 300})
    r = client.get("/api/closing/export", headers=login(client, SAM))
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    assert "Summary" in wb.sheetnames and "Profit" not in wb.sheetnames  # a supervisor's copy
    day = client.get("/api/closing", headers=owner).json()["day"]
    r = client.get(f"/api/reports/export?start={day}&end={day}", headers=owner)
    wb = load_workbook(io.BytesIO(r.content))
    assert {"Profit", "Items", "Sections", "Debtors"} <= set(wb.sheetnames)
    assert wb["Items"]["A2"].value == "Tusker" and wb["Items"]["C2"].value == 300
    assert client.get(f"/api/reports/export?start={day}&end={day}", headers=login(client, SAM)).status_code == 403


def test_closing_is_emailed(config, printers, business, menu, owner, client, db):
    mail_config = config.model_copy(update={"smtp_host": "smtp.example.com", "smtp_from": "faida@example.com"})
    app = create_app(mail_config, printers=printers, start_sync=False)
    with TestClient(app) as c:
        headers = login(c, "2580")
        c.patch("/api/settings", headers=headers, json={"report_emails": ["owner@example.com"]})
        day = c.get("/api/closing", headers=headers).json()["day"]
        assert c.post("/api/closing", json={"day": day}, headers=headers).status_code == 201
    (mail,) = db.scalars(select(EmailOutbox)).all()
    assert mail.to == ["owner@example.com"] and mail.attachment_name == f"closing-{day}.xlsx"
    assert "Test Bar: daily closing" in mail.subject

    sent = []
    assert send_pending(app.state.sessionmaker, mail_config, lambda cfg, msg: sent.append(msg)) == 1
    assert sent[0]["To"] == "owner@example.com"
    assert send_pending(app.state.sessionmaker, mail_config, lambda cfg, msg: sent.append(msg)) == 0
    assert client.patch("/api/settings", headers=owner, json={"report_emails": ["not-an-email"]}).status_code == 422


PNG = "data:image/png;base64," + base64.b64encode(
    bytes.fromhex("89504e470d0a1a0a0000000d4948445200000001000000010806000000"
                  "1f15c4890000000d49444154789c6360000002000154a24f5f0000000049454e44ae426082")).decode()


def test_menu_pictures(client, owner, people, m):
    item = m["items"]["Tusker"]
    r = client.post(f"/api/menu/items/{item['id']}/image", headers=owner, json={"data_url": PNG})
    assert r.status_code == 200
    url = r.json()["image"]
    assert client.get(url).status_code == 200
    assert next(i for c in client.get("/api/menu", headers=owner).json()["categories"] for i in c["items"]
                if i["name"] == "Tusker")["image"] == url
    assert client.post(f"/api/menu/items/{item['id']}/image", headers=owner,
                       json={"data_url": "data:text/plain;base64,aGk="}).status_code == 422
    assert client.post(f"/api/menu/items/{item['id']}/image", headers=login(client, SAM),
                       json={"data_url": PNG}).status_code == 403
    assert client.get("/api/menu/images/../../etc/passwd").status_code == 404


def test_owner_menu_editing(client, owner, m):
    manage = client.get("/api/menu/manage", headers=owner).json()
    assert {g["name"] for g in manage["modifier_groups"]} == {"Mixer", "Side", "How cooked"}
    r = client.post("/api/menu/modifier-groups", headers=owner, json={"name": "Sauce", "min_select": 0, "max_select": 2})
    gid = r.json()["id"]
    client.post(f"/api/menu/modifier-groups/{gid}/options", headers=owner, json={"label": "Pili pili", "price": 20})
    chips = m["items"]["Chips"]
    assert client.put(f"/api/menu/items/{chips['id']}/modifier-groups", headers=owner, json={"group_ids": [gid]}).status_code == 204
    item = next(i for c in client.get("/api/menu", headers=owner).json()["categories"] for i in c["items"] if i["name"] == "Chips")
    assert item["modifier_groups"][0]["options"][0]["label"] == "Pili pili"
    combo = m["items"]["Chicken & chips + Coke"]
    assert client.put(f"/api/menu/items/{combo['id']}/combo", headers=owner,
                      json={"parts": [{"sell_option_id": m["opts"]["Chicken & chips + Coke"]}]}).status_code == 422


def test_split_moves_a_combo_with_its_parts(client, owner, m):
    bill = send(client, owner, tab(client, owner), (m["opts"]["Chicken & chips + Coke"], 1, []), (m["opts"]["Coke"], 1, [])).json()
    parent = next(l for l in bill["lines"] if l["item_name"] == "Chicken & chips + Coke")
    new = client.post(f"/api/bills/{bill['id']}/split", headers=owner, json={"line_ids": [parent["id"]]}).json()
    assert new["total"] == 750 and len(new["lines"]) == 3


def test_sample_pictures_load_by_item_name(db, business, menu, tmp_path):
    from pathlib import Path

    from app.menu.importer import load_images, picture_name
    from app.menu.models import MenuItem

    assert picture_name("Nyama choma 1/2 kg") == "nyama-choma-1-2-kg"
    sample = Path(__file__).parents[2] / "docs" / "sample-images"
    count = load_images(db, business.id, str(sample), tmp_path)
    tusker = db.scalar(select(MenuItem).where(MenuItem.name == "Tusker"))
    assert count >= 10 and tusker.image and (tmp_path / tusker.image).is_file()
    assert load_images(db, business.id, str(sample), tmp_path) == 0  # pictures already set are kept
