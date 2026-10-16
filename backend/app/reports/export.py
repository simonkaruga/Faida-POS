"""Excel copies of the daily closing and of reports, for the owner's accountant or files."""

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="13234A")
WHITE = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
MONEY = "#,##0"


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list[Any]], money_cols: tuple[int, ...] = ()):
    ws = wb.create_sheet(title[:31])
    ws.append(header)
    for cell in ws[1]:
        cell.fill, cell.font = HEAD, WHITE
    for row in rows:
        ws.append(row)
    for col in money_cols:
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.number_format = MONEY
    for i, _ in enumerate(header, start=1):
        width = max(len(str(c.value or "")) for c in ws[get_column_letter(i)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"
    return ws


def _pairs(wb: Workbook, title: str, pairs: list[tuple[str, Any]]):
    ws = _sheet(wb, title, ["", "KSh"], [[k, v] for k, v in pairs], money_cols=(2,))
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value).startswith(("=", "TOTAL", "REAL")):
            for c in row:
                c.font = BOLD
    return ws


def _sales_pairs(s: dict) -> list[tuple[str, Any]]:
    return [("Sales", s["total"]), ("Cash", s["cash"]), ("M-Pesa", s["mpesa"]), ("Card", s.get("card", 0)),
            ("Credit", s["credit"]), ("Unpaid (waiter responsible)", s["unpaid"]), ("Discounts given", s["discounts"]),
            ("M-Pesa tips (not sales)", s["tips"]), ("Bills", s["bills"]), ("Average bill", s["average_bill"])]


def _profit_pairs(p: dict) -> list[tuple[str, Any]]:
    rows = [("Sales", p["sales"]), ("- Cost of goods sold", p["cost_of_sales"]), ("= Gross profit", p["gross_profit"]),
            ("- Waste", p["waste"]), ("+/- Stock-take difference", p["variance"])]
    rows += [(f"- {k.capitalize()}", v) for k, v in p["expenses_by_category"].items()]
    return rows + [("REAL PROFIT", p["real_profit"])]


def _finish(wb: Workbook) -> bytes:
    del wb["Sheet"]
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 110
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(vertical="top")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def closing_workbook(business: str, report: dict) -> bytes:
    wb = Workbook()
    _pairs(wb, "Summary", [("Business", business), ("Business day", report["day"])] + _sales_pairs(report["sales"])
           + [("Credit collected", report["credit_collected"]["total"]),
              ("Cash box should hold", report["cash_box"]["expected_now"])])
    _sheet(wb, "Hand-overs", ["Staff", "Expected", "Received", "Short (-) / over (+)", "Received by", "Signed"],
           [[w["staff"], w["expected"], w["received"], w["difference"], w["receiver"],
             "not present" if w["agrees"] is None else "agrees" if w["agrees"] else "disagrees"] for w in report["waiters"]],
           money_cols=(2, 3, 4))
    if report.get("sections"):
        _sheet(wb, "Sections", ["Section", "Sales"], [[r["section"], r["sales"]] for r in report["sections"]], money_cols=(2,))
    _sheet(wb, "Voids", ["Bill", "Item", "Price", "Reason", "Asked by", "Approved by"],
           [[v["bill"], v["item"], v["price"], v["reason"], v["asked_by"], v["approved_by"]] for v in report["voids"]],
           money_cols=(3,))
    _sheet(wb, "Discounts", ["Bill", "Amount", "Reason", "Asked by", "Approved by"],
           [[d["bill"], d["amount"], d["reason"], d["asked_by"], d["approved_by"]] for d in report["discounts"]], money_cols=(2,))
    if report.get("comps"):
        _sheet(wb, "Free items", ["Bill", "Item", "Reason", "Menu price", "Asked by", "Approved by"],
               [[c["bill"], c["item"], c["reason"], c["value"], c["asked_by"], c["approved_by"]] for c in report["comps"]["items"]],
               money_cols=(4,))
    _sheet(wb, "Unpaid bills", ["Bill", "Place", "Waiter", "Amount"],
           [[u["bill"], u["place"], u["waiter"], u["amount"]] for u in report["unpaid_bills"]], money_cols=(4,))
    _sheet(wb, "Pay-outs", ["What for", "By", "Amount"],
           [[p["note"], p["by"], p["amount"]] for p in report["cash_box"]["payouts"]], money_cols=(3,))
    _sheet(wb, "Stock", ["Item", "Status", "Level", "Value (KSh)"],
           [[v["item"], v["variance"], "", v["value"]] for v in report["stock_variance"]]
           + [[s["item"], "below zero", s["level"], None] for s in report["stock"]["check"]]
           + [[s["item"], "low", s["level"], None] for s in report["stock"]["low"]], money_cols=(4,))
    if "profit" in report:
        _pairs(wb, "Profit", _profit_pairs(report["profit"]))
    return _finish(wb)


def report_workbook(business: str, r: dict) -> bytes:
    wb = Workbook()
    b = r["breakdown"]
    _pairs(wb, "Summary", [("Business", business), ("From", str(r["start"])), ("To", str(r["end"]))]
           + _sales_pairs(r["sales"]) + [("Credit collected", r["credit_collected"]["total"])])
    _pairs(wb, "Profit", _profit_pairs(r["profit"]))
    _sheet(wb, "Items", ["Item", "Quantity", "Sales"], [[i["item"], i["quantity"], i["sales"]] for i in b["by_item"]], money_cols=(3,))
    _sheet(wb, "Categories", ["Category", "Sales"], [[c["category"], c["sales"]] for c in b["by_category"]], money_cols=(2,))
    _sheet(wb, "Waiters", ["Waiter", "Bills", "Sales"], [[w["waiter"], w["bills"], w["sales"]] for w in b["by_waiter"]], money_cols=(3,))
    _sheet(wb, "Sections", ["Section", "Sales"], [[s["section"], s["sales"]] for s in b.get("by_section", [])], money_cols=(2,))
    _sheet(wb, "Hours", ["Hour", "Sales"], [[f"{h['hour']:02d}:00", h["sales"]] for h in b["by_hour"]], money_cols=(2,))
    _sheet(wb, "Expenses", ["Category", "Amount"], [[k, v] for k, v in r["expenses"]["by_category"].items()], money_cols=(2,))
    _sheet(wb, "Shortages", ["Staff", "Shifts", "Times short", "Short", "Times over", "Over"],
           [[s["staff"], s["shifts"], s["times_short"], s["short"], s["times_over"], s["over"]] for s in r["shortages"]],
           money_cols=(4, 6))
    _sheet(wb, "Free items", ["Bill", "Item", "Reason", "Menu price"],
           [[c["bill"], c["item"], c["reason"], c["value"]] for c in r["comps"]["items"]], money_cols=(4,))
    _sheet(wb, "Stock differences", ["Day", "Item", "Difference", "Value"],
           [[v["day"], v["item"], v["variance"], v["value"]] for v in r["variance"]], money_cols=(4,))
    _sheet(wb, "Debtors", ["Name", "Phone", "Owes", "Limit", "Since"],
           [[d["name"], d["phone"], d["owed"], d["limit"], (d["since"] or "")[:10]] for d in r["debtors"]], money_cols=(3, 4))
    return _finish(wb)
